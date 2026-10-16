"""Independent cell/style/type oracles for the bounded mapping subset."""

import io
import json
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable, Mapping
from dataclasses import replace
from hashlib import sha256
from typing import cast

import openpyxl
import pytest
from craftreport_gridcraft.app_values import BrowserTarget
from craftreport_gridcraft.codecs import support_descriptor
from craftreport_gridcraft.values import PROFILE, AppIssue, MappingPlan, OwnedLayout
from craftreport_gridcraft.workbook import (
    bootstrap,
    validate_mapping,
)
from profile_cases import boundary_plans, candidate_run, drawing_signature

from craftreport.adapters.report_codecs import prepared_identifier
from craftreport.contracts.pins import VersionPin
from craftreport.policies.values import QualityEvaluation
from craftreport.reporting.prepared import PreparedReport
from craftreport.themes.values import ValueFormat


@pytest.mark.parametrize("column,limit", [(0, 16), (2, 32)])
@pytest.mark.parametrize("kind", ["ascii", "unicode", "escaped"])
@pytest.mark.parametrize("over", [False, True])
def test_category_and_literal_limits_are_independent(
    column: int,
    limit: int,
    kind: str,
    over: bool,
    prepared: PreparedReport,
    quality_factory: Callable[[PreparedReport], QualityEvaluation],
    layout: OwnedLayout,
) -> None:
    text = {
        "ascii": "W" * (limit - 2),
        "unicode": "é" * ((limit - 2) // 2),
        "escaped": '"\\' + "W" * (limit - 6),
    }[kind] + ("x" if over else "")
    assert len(json.dumps(text, ensure_ascii=False).encode()) == limit + int(over)
    table, chart = prepared.tables[0], prepared.charts[0]
    row = (text, 10, "plain") if column == 0 else ("Alpha", 10, text)
    report = replace(
        prepared,
        tables=(replace(table, rows=(row, *table.rows[1:])),),
        charts=(
            replace(chart, categories=(text, *chart.categories[1:])) if column == 0 else chart,
        ),
    )
    report = replace(report, content_digest=prepared_identifier(report))
    bound = replace(layout, run=candidate_run(report))
    result = validate_mapping(report, quality_factory(report), bound)
    if over:
        assert result == AppIssue("limit", ("tables", 0, "rows", 0, column))
    else:
        assert isinstance(result, MappingPlan)
        assert (
            next(c.value for c in result.cells if c.address == ("A6" if column == 0 else "C6"))
            == text
        )


@pytest.fixture
def layout(prepared: PreparedReport, browser_target: BrowserTarget) -> OwnedLayout:
    return OwnedLayout(browser_target, candidate_run(prepared))


def test_exact_mapping_and_text_safe_bootstrap(
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
) -> None:
    plan = validate_mapping(prepared, quality, layout)
    assert isinstance(plan, MappingPlan)
    cells = {c.address: c.value for c in plan.cells}
    expected = {
        "A1": "DRAFT",
        "A2": "Supplied series",
        "A3": quality.validation_summary,
        "A5": "Category",
        "B5": "Value",
        "C5": "Literal",
        "A6": "Alpha",
        "B7": 15,
        "C6": "=1+1",
        "C7": "https://example.invalid/",
        "C8": "plain",
        "C9": "check",
        "G7": None,
        "F9": "Delta",
        "G9": 20,
        "J1": "document",
        "K1": "document-marker",
        "K2": "run-1",
        "K11": sha256(quality.validation_summary.encode()).hexdigest(),
    }
    assert {a: cells[a] for a in expected} == expected
    data = bootstrap(plan, "document-marker")
    assert isinstance(data, bytes)
    book = openpyxl.load_workbook(io.BytesIO(data))
    try:
        sheet = book["Report"]
        assert {a: sheet[a].value for a in expected} == expected
        assert all(sheet[a].data_type == "s" for a in ("C6", "C7", "C8", "C9"))
        assert sheet["C7"].hyperlink is None and sheet["B7"].data_type == "n"
        assert sheet["B7"].number_format == "0.00;-0.00"
        assert sheet["G7"].data_type == "n"
        assert sheet["M1"].value == "OUTSIDE-UNCHANGED"
        assert sheet.sheet_view.showGridLines is False
        assert getattr(sheet.row_dimensions[5], "height") == 24
        assert sheet["A1"].font.name == "Arial" and sheet["A1"].font.bold is True
        assert sheet["A3"].alignment.wrap_text is True
    finally:
        book.close()
    assert isinstance(bootstrap(plan, "other-marker"), AppIssue)


@pytest.mark.parametrize("value", [-1, 10000])
def test_reject_numeric_bounds_before_bootstrap(
    value: int,
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
) -> None:
    table = prepared.tables[0]
    bad = replace(
        prepared, tables=(replace(table, rows=(("Alpha", value, "=1+1"), *table.rows[1:])),)
    )
    issue = validate_mapping(bad, quality, layout)
    assert isinstance(issue, AppIssue) and issue.code == "limit"
    assert issue.location == ("tables", 0, "rows", 0, 1)


def test_schema_theme_shape_quality_and_pin_fail_closed(
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
) -> None:
    table, chart = prepared.tables[0], prepared.charts[0]
    candidates = (
        replace(prepared, title="Other title"),
        replace(prepared, tables=(replace(table, rows=table.rows[:3]),)),
        replace(prepared, charts=(replace(chart, categories=("Other", *chart.categories[1:])),)),
        replace(
            prepared,
            tables=(
                replace(
                    table,
                    columns=(replace(table.columns[0], display_label="Other"), *table.columns[1:]),
                ),
            ),
        ),
        replace(
            prepared,
            presentation=replace(
                prepared.presentation, formats=(ValueFormat("text"), ValueFormat("int64", 3))
            ),
        ),
        replace(prepared, required_capabilities=()),
    )
    assert all(isinstance(validate_mapping(p, quality, layout), AppIssue) for p in candidates)
    assert isinstance(
        validate_mapping(prepared, replace(quality, draft_eligible=False), layout), AppIssue
    )
    changed = replace(
        quality.findings[0], scope=replace(quality.findings[0].scope, content_digest="0" * 64)
    )
    assert isinstance(
        validate_mapping(
            prepared, replace(quality, findings=(changed, *quality.findings[1:])), layout
        ),
        AppIssue,
    )
    assert isinstance(
        validate_mapping(prepared, replace(quality, validation_summary="x" * 1024), layout),
        AppIssue,
    )
    assert isinstance(
        validate_mapping(
            prepared,
            quality,
            replace(
                layout,
                target=replace(layout.target, profile_pin=VersionPin("other", "1", "a" * 64)),
            ),
        ),
        AppIssue,
    )
    assert isinstance(
        validate_mapping(
            prepared, quality, replace(layout, run=replace(layout.run, prepared_digest="0" * 64))
        ),
        AppIssue,
    )


@pytest.mark.parametrize("value", [0, 1, 9999])
def test_candidate_numeric_domain_is_exact(
    value: int,
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
    quality_factory: Callable[[PreparedReport], QualityEvaluation],
) -> None:
    table = prepared.tables[0]
    candidate = replace(
        prepared, tables=(replace(table, rows=(("Alpha", value, "=1+1"), *table.rows[1:])),)
    )
    candidate = replace(candidate, content_digest=prepared_identifier(candidate))
    layout = replace(layout, run=candidate_run(candidate))
    quality = quality_factory(candidate)
    plan = validate_mapping(candidate, quality, layout)
    assert isinstance(plan, MappingPlan)
    assert next(c.value for c in plan.cells if c.address == "B6") == value
    assert PROFILE.operation_status == "untested"
    assert validate_mapping(candidate, quality, layout) == plan


@pytest.mark.parametrize(
    "text", ["\n", "\r", "\t", "\x00", "\x85", "\u200b", "\u2028", "\u2029", "\u202e"]
)
def test_unproved_controls_reject_without_normalizing(
    text: str,
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
) -> None:
    table = prepared.tables[0]
    changed = replace(
        prepared, tables=(replace(table, rows=(("Alpha", 10, text), *table.rows[1:])),)
    )
    assert validate_mapping(changed, quality, layout) == AppIssue(
        "untested", ("tables", 0, "rows", 0, 2)
    )


@pytest.mark.parametrize("text", ["x" * 30, "é" * 15, "=SUM(A1)", '"\\quoted'])
def test_literal_utf8_and_escaped_boundary_preserves_exact_text(
    text: str,
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
    quality_factory: Callable[[PreparedReport], QualityEvaluation],
) -> None:
    assert len(json.dumps(text, ensure_ascii=False).encode()) <= 32
    table = prepared.tables[0]
    changed = replace(
        prepared, tables=(replace(table, rows=(("Alpha", 10, text), *table.rows[1:])),)
    )
    changed = replace(changed, content_digest=prepared_identifier(changed))
    layout = replace(layout, run=candidate_run(changed))
    plan = validate_mapping(changed, quality_factory(changed), layout)
    assert isinstance(plan, MappingPlan)
    data = bootstrap(plan, layout.document_marker)
    assert isinstance(data, bytes)
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    try:
        assert book["Report"]["C6"].value == text
        assert book["Report"]["C6"].data_type == "s"
    finally:
        book.close()


def test_independent_chart_sources_and_all_applied_cell_styles(
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
) -> None:
    plan = validate_mapping(prepared, quality, layout)
    assert isinstance(plan, MappingPlan)
    data = bootstrap(plan, layout.document_marker)
    assert isinstance(data, bytes)
    book = openpyxl.load_workbook(io.BytesIO(data))
    try:
        sheet = book["Report"]
        assert [sheet[f"B{i}"].value for i in range(6, 10)] == [10, 15, 30, 20]
        assert [sheet[f"G{i}"].value for i in range(6, 10)] == [10, None, 30, 20]
        bold = {"A1", "A2", "A5", "B5", "C5", "F5", "G5", *(f"J{i}" for i in range(1, 13))}
        filled = bold - {"A2"}
        for address in (c.address for c in plan.cells):
            cell = sheet[address]
            assert cell.font.name == "Arial" and cell.font.sz == 11
            assert bool(cell.font.bold) == (address in bold)
            assert cell.font.italic is False
            assert cell.font.color is not None and cell.font.color.rgb == "FF000000"
            assert cell.alignment.horizontal == "left" and cell.alignment.wrap_text is True
            assert cell.fill.fgColor.rgb == ("FFDCE6F1" if address in filled else "00000000")
        assert "A3:H3" in sheet.merged_cells
    finally:
        book.close()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        chart = ET.fromstring(archive.read("xl/charts/chart1.xml"))
        assert b"SAFE" in archive.read("xl/drawings/drawing1.xml")
    ns = {
        "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    assert [x.text for x in chart.findall(".//c:f", ns)] == [
        "Report!$G$5",
        "Report!$F$6:$F$9",
        "Report!$G$6:$G$9",
    ]
    line = chart.find(".//c:ser/c:spPr/a:ln", ns)
    assert line is not None and line.get("w") == "28575"
    color = line.find("a:solidFill/a:srgbClr", ns)
    assert color is not None and color.get("val") == "4472C4"
    assert chart.find(".//c:legend", ns) is None
    assert chart.find(".//c:majorGridlines", ns) is None


def test_approved_width_is_exact_and_boundary_plans_keep_paired_pins(
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
    quality_factory: Callable[[PreparedReport], QualityEvaluation],
) -> None:
    plan = validate_mapping(prepared, quality, layout)
    assert isinstance(plan, MappingPlan)
    assert prepared.presentation.chart.line_width == 2.25
    for width in (2, 2.5):
        altered = replace(
            prepared,
            presentation=replace(
                prepared.presentation, chart=replace(prepared.presentation.chart, line_width=width)
            ),
        )
        assert validate_mapping(altered, quality, layout) == AppIssue("untested", ("presentation",))
    for candidate in boundary_plans(plan, quality_factory):
        data = bootstrap(candidate, layout.document_marker)
        assert isinstance(data, bytes)
        drawing_signature(data)
        assert candidate.layout.prepared_digest == candidate.prepared.content_digest
        assert all(
            f.scope.content_digest == candidate.prepared.content_digest
            for f in candidate.quality.findings
        )


def test_exact_run_identity_replaces_test_marker(
    prepared: PreparedReport, quality: QualityEvaluation, layout: OwnedLayout
) -> None:
    bound = replace(layout, run=replace(layout.run, run_id="retained-run-42"))
    plan = validate_mapping(prepared, quality, bound)
    assert isinstance(plan, MappingPlan)
    assert next(c.value for c in plan.cells if c.address == "K2") == "retained-run-42"


@pytest.mark.parametrize(
    "field,value",
    [
        ("parameter_digest", "0" * 64),
        ("validation_scope", "0" * 64),
        ("resolved_parameters", (("unexpected", 1),)),
    ],
)
def test_run_parameter_and_scope_mismatch_rejects(
    prepared: PreparedReport,
    quality: QualityEvaluation,
    layout: OwnedLayout,
    field: str,
    value: object,
) -> None:
    bound = replace(layout, run=replace(layout.run, **{field: value}))
    assert validate_mapping(prepared, quality, bound) == AppIssue("conflict", ("run",))


def test_forged_content_and_changed_build_fail_before_bootstrap(
    prepared: PreparedReport, quality: QualityEvaluation, layout: OwnedLayout
) -> None:
    table = prepared.tables[0]
    forged = replace(
        prepared, tables=(replace(table, rows=(("Alpha", 11, "=1+1"), *table.rows[1:])),)
    )
    assert validate_mapping(forged, quality, layout) == AppIssue(
        "conflict", ("prepared", "content_digest")
    )
    changed = replace(layout, target=replace(layout.target, wasm_digest="0" * 64))
    assert validate_mapping(prepared, quality, changed) == AppIssue("untested", ("profile",))


def test_mapping_support_does_not_activate_runtime_operations() -> None:
    descriptor = support_descriptor()
    assert descriptor["mapping_status"] == "supported"
    assert descriptor["operation_status"] == "untested"
    assert descriptor["headed_status"] == "MANUAL_VERIFY_REQUIRED"
    properties = cast(tuple[Mapping[str, object], ...], descriptor["properties"])
    excluded = {
        "format.negative_style",
        "format.date_pattern",
        "table.typed_null",
        "table.null_marker",
        "table.empty_headers",
        "table.empty_note",
    }
    assert len(properties) == 24
    assert {p["name"] for p in properties if p["status"] == "untested"} == excluded
    assert sum(p["status"] == "supported" for p in properties) == 18
