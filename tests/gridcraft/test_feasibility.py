"""Disposable actual released-app route probe; no production capability claim."""

import asyncio
import hashlib
import io
import os
import platform
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Awaitable, Callable
from pathlib import Path
from typing import cast
from uuid import uuid4

import openpyxl
import pytest
import test_browser as host
import xlsxwriter  # pyright: ignore[reportMissingTypeStubs] -- pinned library lacks py.typed
from PIL import Image
from playwright.async_api import Page

from scripts.browser_verification import browser, deployment
from scripts.browser_verification.config import OwnedDeployment, RunConfig
from scripts.browser_verification.interactions import click, copy_text, frame, key, selected

RECEIVER = """<script>
window.saveLane={armed:false,count:0,failed:false,candidate:null};
addEventListener('message',e=>{
 if(e.origin!==location.origin || e.source!==app.contentWindow) return;
 const d=e.data;
 if(!d || d.type!=='gridcraft:save') return;
 const s=window.saveLane;
 s.count++;
 if(!s.armed || s.candidate || !(d.bytes instanceof ArrayBuffer)
    || d.bytes.byteLength>4194304 || typeof d.name!=='string') {s.failed=true;return;}
 s.candidate={name:d.name,bytes:Array.from(new Uint8Array(d.bytes))};
 s.armed=false;
});</script>"""


def fixture(marker: str) -> bytes:
    output = io.BytesIO()
    with xlsxwriter.Workbook(
        output, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}
    ) as book:
        # Only third-party optional factory parameters lack annotations; values remain typed.
        sheet = book.add_worksheet("Report")  # pyright: ignore[reportUnknownMemberType]
        bold = book.add_format({"bold": True, "bg_color": "#DCE6F1"})  # pyright: ignore[reportUnknownMemberType]
        numeric = book.add_format({"num_format": "0.00;-0.00"})  # pyright: ignore[reportUnknownMemberType]
        assert sheet.write_string(0, 0, "DRAFT", bold) == 0
        assert sheet.write_string(0, 9, marker) == 0
        assert sheet.write_string(1, 9, "run-1") == 0
        assert sheet.write_string(0, 12, "OUTSIDE-UNCHANGED") == 0
        for address, value in {
            "A5": "Category",
            "B5": "Value",
            "C5": "Literal",
            "A6": "Alpha",
            "A7": "Beta",
            "A8": "Gamma",
            "C6": "=1+1",
            "F5": "Category",
            "G5": "Series",
            "F6": "Alpha",
            "F7": "Beta",
            "F8": "Gamma",
            "A9": "Delta",
            "F9": "Delta",
        }.items():
            assert sheet.write_string(int(address[1:]) - 1, ord(address[0]) - 65, value) == 0
        for address, value in {"B6": 10, "B8": 30, "G6": 10, "G8": 30, "B9": 20, "G9": 20}.items():
            assert (
                sheet.write_number(int(address[1:]) - 1, ord(address[0]) - 65, value, numeric) == 0
            )
        assert sheet.write_blank(6, 1, None, numeric) == 0
        assert sheet.write_blank(6, 6, None, numeric) == 0
        chart = book.add_chart({"type": "line"})
        assert chart is not None
        chart.add_series(
            {
                "name": "Report!$G$5",
                "categories": "Report!$F$6:$F$9",
                "values": "Report!$G$6:$G$9",
                "line": {"color": "#4472C4"},
            }
        )
        chart.set_title({"name": "Supplied series"})
        chart.show_blanks_as("gap")
        assert sheet.insert_chart(17, 0, chart, {"x_scale": 0.8, "y_scale": 0.8}) == 0
        assert sheet.insert_textbox(2, 12, "SAFE", {"width": 120, "height": 35}) == 0
    return output.getvalue()


async def cell(page: Page, address: str) -> dict[str, str]:
    await click(page, 40, 179)
    await key(page, "Control+a")
    await page.keyboard.insert_text(address)
    await frame(page)
    await key(page, "Enter")
    async with asyncio.timeout(5):
        while not selected(await page.screenshot(), address):
            await asyncio.sleep(0.05)
    await click(page, 40, 179)
    await key(page, "Control+a")
    async with asyncio.timeout(5):
        reference, rn = await copy_text(page)
    assert reference == address, (address, reference)
    await key(page, "Escape")
    child = page.frames[1]
    await child.locator("#gridcraft_canvas").focus()
    async with asyncio.timeout(5):
        value, vn = await copy_text(page)
    return {"reference": reference, "value": value, "reference_nonce": rn, "value_nonce": vn}


async def save(page: Page, directory: Path, name: str) -> bytes:
    await page.evaluate(
        "() => {if(window.saveLane.failed) throw Error('lane failed');"
        "window.saveLane={armed:true,count:0,failed:false,candidate:null};}"
    )
    await key(page, "Control+s")
    await page.wait_for_function(
        "window.saveLane.candidate!==null || window.saveLane.failed", timeout=10000
    )
    lane: dict[str, object] = await page.evaluate("window.saveLane")
    host.write_json(directory / (name + "-lane.json"), lane)
    assert lane["failed"] is False and lane["count"] == 1
    raw: list[int] = await page.evaluate("window.saveLane.candidate.bytes")
    output = bytes(raw)
    path = directory / (name + ".xlsx")
    with path.open("xb") as stream:
        stream.write(output)
        stream.flush()
        os.fsync(stream.fileno())
    assert path.read_bytes() == output
    return output


def inspect(data: bytes, marker: str, category: str, number: int, run: str) -> dict[str, object]:
    assert len(data) <= 4 * 1024 * 1024
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert len(archive.infolist()) <= 256
        assert sum(x.file_size for x in archive.infolist()) <= 32 * 1024 * 1024
        parts = {p: archive.read(p) for p in archive.namelist()}
        assert all(b"<!DOCTYPE" not in v and b"<!ENTITY" not in v for v in parts.values())
        chart = ET.fromstring(parts["xl/charts/chart1.xml"])
        drawings = ET.fromstring(parts["xl/drawings/drawing1.xml"])
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    assert [x.text for x in chart.findall(".//c:f", ns)] == [
        "Report!$G$5",
        "Report!$F$6:$F$9",
        "Report!$G$6:$G$9",
    ]
    assert [x.text for x in chart.findall(".//c:cat//c:v", ns)] == [
        "Alpha",
        "Beta",
        "Gamma",
        category,
    ]
    assert [x.text for x in chart.findall(".//c:val//c:v", ns)] == ["10", "30", str(number)]
    gap = chart.find(".//c:dispBlanksAs", ns)
    assert gap is not None and gap.get("val") == "gap"
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    expected = {
        "J1": marker,
        "J2": run,
        "A9": category,
        "F9": category,
        "B9": number,
        "G9": number,
        "B7": None,
        "G7": None,
        "C6": "=1+1",
        "M1": "OUTSIDE-UNCHANGED",
    }
    try:
        sheet = book["Report"]
        assert {a: sheet[a].value for a in expected} == expected
        assert sheet["C6"].data_type == "s"
        assert sheet["A1"].font.bold and sheet["A1"].fill.fgColor.rgb == "FFDCE6F1"
        assert sheet["B6"].number_format == "0.00;-0.00"
    finally:
        book.close()
    outside = ET.tostring(list(drawings)[1]).decode()
    assert "SAFE" in outside
    return {"sha256": hashlib.sha256(data).hexdigest(), "values": expected, "outside": outside}


async def update(
    page: Page, marker: str, before: tuple[str, int, str], after: tuple[str, int, str]
) -> None:
    assert (await cell(page, "J1"))["value"] == marker
    addresses = ("A9", "F9", "B9", "G9", "J2")
    old = (before[0], before[0], str(before[1]), str(before[1]), before[2])
    new = (after[0], after[0], str(after[1]), str(after[1]), after[2])
    for address, value in zip(addresses, old, strict=True):
        observed = (await cell(page, address))["value"]
        assert observed == value or observed == value + ".00", (address, observed, value)
    for address, value in zip(addresses, new, strict=True):
        await cell(page, address)
        await click(page, 300, 179)
        await key(page, "Control+a")
        await page.keyboard.insert_text(value)
        await frame(page)
        await key(page, "Enter")


async def bootstrap(page: Page, data: bytes, directory: Path) -> None:
    await page.context.grant_permissions(["clipboard-read", "clipboard-write"], origin=page.url)
    await page.evaluate(
        "data => app.contentWindow.postMessage({type:'gridcraft:open',"
        "name:'reference.xlsx',bytes:new Uint8Array(data)},location.origin)",
        list(data),
    )
    await asyncio.sleep(2)
    await page.screenshot(path=str(directory / "opened.png"))


async def observe_cells(
    page: Page, marker: str, expected: tuple[str, int, str], directory: Path, name: str
) -> None:
    values = {
        "J1": marker,
        "J2": expected[2],
        "A9": expected[0],
        "F9": expected[0],
        "B9": str(expected[1]) + ".00",
        "G9": str(expected[1]) + ".00",
        "C6": "=1+1",
        "M1": "OUTSIDE-UNCHANGED",
    }
    observed = [await cell(page, a) for a in values]
    host.write_json(directory / (name + "-cells.json"), observed)
    assert {x["reference"]: x["value"] for x in observed} == values
    await cell(page, "A1")
    await key(page, "Escape")
    await page.screenshot(path=str(directory / (name + "-visual.png")))
    visual(directory / (name + "-visual.png"), expected[1])


def visual(path: Path, value: int) -> None:
    with Image.open(path) as image:
        pixels = image.convert("RGB")
        assert pixels.size == (1600, 1100)
        blue = [
            (x, y)
            for x in range(60, 285)
            for y in range(515, 651)
            if (lambda c: c[2] > c[0] * 1.5 and c[2] > c[1] * 1.3)(
                cast(tuple[int, int, int], pixels.getpixel((x, y)))
            )
        ]
        assert len(blue) > 50
        assert min(x for x, _ in blue) >= 197  # No segment over the null gap.
        assert any(abs(x - 199) < 3 and abs(y - 536) < 3 for x, y in blue)
        assert any(abs(x - 255) < 3 and abs(y - (650 - value * 3.8)) < 4 for x, y in blue)


async def journey(page: Page, marker: str, directory: Path) -> bytes:
    previous = ("Delta", 20, "run-1")
    outside = None
    final = b""
    for index, expected in enumerate((previous, ("Epsilon", 5, "run-2"), ("Zeta", 25, "run-3"))):
        if index:
            await update(page, marker, previous, expected)
        await observe_cells(page, marker, expected, directory, f"stage-{index}")
        final = await save(page, directory, f"stage-{index}")
        result = inspect(final, marker, *expected)
        host.write_json(directory / f"stage-{index}-inspection.json", result)
        if outside is not None:
            assert result["outside"] == outside
        outside = result["outside"]
        previous = expected
    return final


def conflict_worksheet(data: bytes) -> tuple[bytes, dict[str, str]]:
    assert b"<!DOCTYPE" not in data and b"<!ENTITY" not in data
    parser = ET.XMLParser(target=ET.TreeBuilder(insert_comments=True, insert_pis=True))
    root = ET.fromstring(data, parser=parser)
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    assert root.tag == "{" + ns["s"] + "}worksheet"
    path = ""
    for tag in ("sheetViews", "sheetView", "selection"):
        path += ("/" if path else "") + "s:" + tag
        assert len(root.findall(path, ns)) == 1, path
    selection = root.find(path, ns)
    assert selection is not None and len(selection) == 0
    assert {"activeCell", "sqref"} <= selection.attrib.keys()
    location = {name: selection.attrib.pop(name) for name in ("activeCell", "sqref")}
    return ET.tostring(root), location


def compare_conflict(before: bytes, after: bytes) -> dict[str, dict[str, str]]:
    with zipfile.ZipFile(io.BytesIO(before)) as left, zipfile.ZipFile(io.BytesIO(after)) as right:
        assert left.namelist() == right.namelist()
        assert len(left.namelist()) == len(set(left.namelist()))
        locations: dict[str, dict[str, str]] = {}
        for name in left.namelist():
            a, b = left.read(name), right.read(name)
            if name == "xl/worksheets/sheet1.xml":
                a, locations["before"] = conflict_worksheet(a)
                b, locations["after"] = conflict_worksheet(b)
            assert a == b, name
        assert locations
    return locations


async def conflict(page: Page, marker: str, directory: Path) -> None:
    await cell(page, "A9")
    await click(page, 300, 179)
    await key(page, "Control+a")
    await page.keyboard.insert_text("EXTERNAL")
    await frame(page)
    await key(page, "Enter")
    before = await save(page, directory, "external-before")
    with pytest.raises(AssertionError, match="A9"):
        await update(page, marker, ("Zeta", 25, "run-3"), ("Overwritten", 99, "run-4"))
    after = await save(page, directory, "external-after")
    locations = compare_conflict(before, after)
    host.write_json(
        directory / "conflict.json",
        {
            "rejected_before_mutation": True,
            "before": hashlib.sha256(before).hexdigest(),
            "after": hashlib.sha256(after).hexdigest(),
            "selection_view": locations,
        },
    )


async def reopen(
    saved: bytes,
    config: RunConfig,
    lease: deployment.DeploymentLease,
    state: browser.BrowserState,
    marker: str,
    original: Callable[
        [RunConfig, deployment.DeploymentLease, browser.BrowserState], Awaitable[dict[str, object]]
    ],
) -> None:
    child_config = config.model_copy(update={"output_root": config.run_directory / "reopen"})
    child_config.run_directory.mkdir(parents=True)
    child_state = browser.BrowserState(state.registry)
    try:
        fresh = await original(child_config, lease, child_state)
        assert child_state.page
        assert child_state.diagnostics
        child_state.diagnostics.stage = "open"
        await bootstrap(child_state.page, saved, child_config.run_directory)
        await observe_cells(
            child_state.page,
            marker,
            ("Zeta", 25, "run-3"),
            child_config.run_directory,
            "reopened",
        )
        assert child_state.diagnostics
        assert all(not d.fatal for d in child_state.diagnostics.events)
        fresh["opened_sha256"] = hashlib.sha256(saved).hexdigest()
        host.write_json(child_config.run_directory / "reopen.json", fresh)
    finally:
        await browser.close(child_config, child_state)
        if child_state.diagnostics:
            child_state.diagnostics.write(child_config.run_directory / "diagnostics.json")
        host.write_json(
            child_config.run_directory / "video.json",
            child_state.video.model_dump(mode="json") if child_state.video else None,
        )
    assert child_state.video and child_state.video.decoded_frame


@pytest.mark.integration
@pytest.mark.skipif(platform.system() != "Windows", reason="approved Windows profile")
def test_released_same_document_feasibility(monkeypatch: pytest.MonkeyPatch) -> None:
    config = RunConfig(
        deployment=OwnedDeployment(mode="owned-wsl"),
        headed=False,
        output_root=host.ROOT / "artifacts/gridcraft/task-3",
        trusted_release_sha256=host.ARCHIVE,
        trust_source="approved v0.7.0 Design",
    )
    config.run_directory.mkdir(parents=True)
    marker = "task3-" + uuid4().hex
    original = host.observe_bridge

    async def probe(
        config: RunConfig, lease: deployment.DeploymentLease, state: browser.BrowserState
    ) -> dict[str, object]:
        report = await original(config, lease, state)
        assert state.page
        page = state.page
        data = fixture(marker)
        (config.run_directory / "bootstrap.xlsx").write_bytes(data)
        report["test_source_sha256"] = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
        report["marker"] = marker
        report["scope"] = "task-3 minimal real route only; full theme support untested"
        assert state.diagnostics
        state.diagnostics.stage = "open"
        await bootstrap(page, data, config.run_directory)
        host.write_json(config.run_directory / "probe.json", report)
        assert state.diagnostics and all(not d.fatal for d in state.diagnostics.events)
        state.diagnostics.stage = "edit"
        saved = await journey(page, marker, config.run_directory)
        assert all(not d.fatal for d in state.diagnostics.events)
        await reopen(saved, config, lease, state, marker, original)
        state.diagnostics.stage = "edit"
        await conflict(page, marker, config.run_directory)
        assert all(not d.fatal for d in state.diagnostics.events)
        report["final_sha256"] = hashlib.sha256(saved).hexdigest()
        report["same_document_updates"] = 2
        report["support"] = {
            "shape": "4 rows,3 columns,1 line series; fixed binding",
            "proved": [
                "literal text",
                "bounded small integer",
                "null gap",
                "bold",
                "fill",
                "0.00;-0.00 format",
                "one series color",
                "outside cell/textbox",
            ],
            "untested": [
                "full theme descriptor",
                "decimal precision",
                "bool",
                "date",
                "UTC datetime",
                "empty input",
                "font",
                "alignment",
                "wrap",
                "width/height",
                "grid style",
                "chart line width",
                "arbitrary styles/shapes",
            ],
            "unsupported": ["requested axis bounds"],
            "headed": "MANUAL_VERIFY_REQUIRED",
        }
        host.write_json(config.run_directory / "result.json", report)
        return report

    monkeypatch.setattr(host, "PARENT", host.PARENT + RECEIVER)
    monkeypatch.setattr(host, "observe_bridge", probe)
    asyncio.run(host.readiness(config, monkeypatch, observer_timeout=180))
